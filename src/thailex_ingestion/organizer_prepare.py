"""Turn organizer spreadsheets and Word files into flat CSV for the organizer adapter.

The organizer adapter reads CSV/JSON/JSONL/XML with a declarative mapping, but
the supplied dictionaries are XLSX/DOCX files whose cells mix several facts
(part-of-speech abbreviations, usage labels, pronunciations and numbered
senses inside one definition). Each profile below only *separates* those parts
into columns. It never rewrites, translates or invents content; the original
row is kept in the `raw` column so the adapter stores it as the raw record.
"""

from __future__ import annotations

import csv
import html
import json
import re
import unicodedata
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

OUTPUT_COLUMNS = (
    "id", "lemma", "pos", "sense_number", "definition", "pronunciations",
    "etymology", "notes", "translations", "raw",
)
LIST_SEPARATOR = "|"

# Part-of-speech abbreviations used by the Royal Institute dictionaries.
POS_ABBREVIATIONS = ("สรรพ", "สัน", "นิ", "วิ", "น", "ก", "ว", "บ", "อ", "ส")
_POS = re.compile(r"^(?P<pos>" + "|".join(POS_ABBREVIATIONS) + r")\.(?=\s|$)\s*")
_LOOSE_POS = re.compile(r"^(?P<pos>น|ว|ก)\.\s*")
_BRACKET = re.compile(r"^\[(?P<text>[^\]]*)\]\s*")
_LABEL = re.compile(r"^\((?P<text>[^()]{1,40})\)\s*")
_HOMOGRAPH = re.compile(r"\s+(?P<number>[๐-๙0-9]+)$")
_NUMBERED = re.compile(r"(?:^|(?<=\s))(?P<number>[๐-๙0-9]{1,2})\.\s+")
_TAG = re.compile(r"<[^>]+>")
_THAI_DIGITS = str.maketrans("๐๑๒๓๔๕๖๗๘๙", "0123456789")


def clean(value: object) -> str:
    text = "" if value is None else str(value)
    if text.strip().lower() == "null":
        return ""
    text = html.unescape(_TAG.sub("", text))
    text = re.sub(r"-\s*\n\s*", "", text)  # a word hyphenated across a line break
    text = re.sub(r"[ \t ]+", " ", text.replace("\r", ""))
    text = re.sub(r"\s*\n\s*", " ", text)
    return unicodedata.normalize("NFC", text).strip()


@dataclass
class Leading:
    pos: str = ""
    labels: list[str] = field(default_factory=list)
    pronunciations: list[str] = field(default_factory=list)
    remainder: str = ""


def split_leading(text: str, *, loose_pos: bool = False) -> Leading:
    """Peel `[pronunciation]`, `(usage label)` and one POS abbreviation off the start."""
    result = Leading()
    rest = text.strip()
    while rest:
        bracket = _BRACKET.match(rest)
        if bracket:
            content = bracket.group("text").strip()
            if "อ่านว่า" in content:
                label, _, reading = content.partition("อ่านว่า")
                if label.strip(" ;,"):
                    result.labels.append(f"({label.strip(' ;,')})")
                if reading.strip():
                    result.pronunciations.append(reading.strip())
            elif content:
                result.pronunciations.append(content)
            rest = rest[bracket.end():]
            continue
        label = _LABEL.match(rest)
        if label:
            result.labels.append(f"({label.group('text').strip()})")
            rest = rest[label.end():]
            continue
        if not result.pos:
            pos = (_LOOSE_POS if loose_pos else _POS).match(rest)
            if pos:
                result.pos = pos.group("pos") + "."
                rest = rest[pos.end():]
                continue
        break
    result.remainder = rest.strip()
    return result


def numbered_segments(text: str) -> list[tuple[str, str]]:
    """Split "๑. ก ๒. ข" into [("1", "ก"), ("2", "ข")]; unnumbered text is one segment."""
    matches = list(_NUMBERED.finditer(text))
    if not matches or matches[0].start() > 0 and text[: matches[0].start()].strip():
        return [("", text.strip())]
    segments = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        segments.append((match.group("number").translate(_THAI_DIGITS), text[match.end():end].strip()))
    return [segment for segment in segments if segment[1]]


def _row(**values: object) -> dict[str, str]:
    row = {column: "" for column in OUTPUT_COLUMNS}
    for key, value in values.items():
        if isinstance(value, (list, tuple)):
            value = LIST_SEPARATOR.join(dict.fromkeys(str(item) for item in value if str(item).strip()))
        row[key] = "" if value is None else str(value)
    return row


def _raw(header: Iterable[str], cells: Iterable[object]) -> str:
    return json.dumps(
        {f"{index}:{name}": clean(value) for index, (name, value) in enumerate(zip(header, cells)) if clean(value)},
        ensure_ascii=False,
    )


# --- Spreadsheet profiles -------------------------------------------------


def royal_2542(header: list[str], rows: Iterable[list[object]]) -> Iterator[dict[str, str]]:
    """Columns: number, K_W (search keys), M_W (headword + homograph number), D_T (entry text).

    `number` is blank for most of the file, so the row position is the stable id.
    One row lacks M_W; its last search key is the headword form.
    """
    for index, cells in enumerate(rows, start=1):
        cells = list(cells) + [""] * (4 - len(cells))
        keys = [key.strip() for key in clean(cells[1]).split(",") if key.strip()]
        headword = clean(cells[2]) or (keys[-1] if keys else "")
        text = clean(cells[3])
        if not headword or not text:
            continue
        homograph = _HOMOGRAPH.search(headword)
        written = headword[: homograph.start()] if homograph else headword
        leading = split_leading(text)
        # "นย-, นยะ" lists spelling variants of one entry; the trailing hyphen
        # marks a combining form, so each variant is looked up without it.
        variants = list(dict.fromkeys(
            part.strip().rstrip("-").strip() for part in written.split(",") if part.strip().rstrip("-").strip()
        )) or [written]
        for variant_number, lemma in enumerate(variants, start=1):
            yield _row(
                id=f"2542-{index:05d}" + (f"-{variant_number}" if len(variants) > 1 else ""),
                lemma=lemma, pos=leading.pos,
                sense_number=homograph.group("number") if homograph else "",
                definition=leading.remainder or text, pronunciations=leading.pronunciations,
                notes=[*leading.labels, *([f"รูปคำในต้นฉบับ: {written}"] if len(variants) > 1 else [])],
                raw=_raw(header, cells),
            )


def royal_2554(header: list[str], rows: Iterable[list[object]]) -> Iterator[dict[str, str]]:
    """One row per definition; rows sharing `number` are sub-senses of one headword sense."""
    labelled = {
        4: "สาขาวิชา", 7: "อ้างอิง", 9: "การใช้", 10: "อ้างอิง", 12: "ระดับภาษา",
        15: "ดู", 20: "รูปเต็ม", 21: "รูปย่อ", 22: "ราชาศัพท์",
    }
    previous_number, part, group_pos = "", 0, ""
    for cells in rows:
        cells = list(cells) + [""] * (23 - len(cells))
        number, lemma, definition_text = clean(cells[0]), clean(cells[1]), clean(cells[6])
        if not number or not lemma:
            continue
        part = part + 1 if number == previous_number else 1
        if number != previous_number:
            group_pos = clean(cells[5])
        previous_number = number
        leading = split_leading(definition_text)
        notes = list(leading.labels)
        if clean(cells[13]):
            notes.append(f"({clean(cells[13])})")
        notes.extend(f"{label}: {clean(cells[index])}" for index, label in labelled.items() if clean(cells[index]))
        etymology = " ; ".join(value for value in (clean(cells[11]), clean(cells[14])) if value)
        yield _row(
            id=f"2554-{number}-{part}", lemma=lemma, pos=clean(cells[5]) or leading.pos or group_pos,
            sense_number=clean(cells[2]), definition=leading.remainder or definition_text,
            pronunciations=[clean(cells[3]), *leading.pronunciations] if part == 1 else leading.pronunciations,
            etymology=etymology, notes=notes, raw=_raw(header, cells),
        )


def royal_2569(header: list[str], rows: Iterable[list[object]]) -> Iterator[dict[str, str]]:
    """Columns: head_word, snumber_sense, pronunciation, part_of_speech, definition."""
    for index, cells in enumerate(rows, start=1):
        cells = list(cells) + [""] * (5 - len(cells))
        lemma, definition_text = clean(cells[0]), clean(cells[4])
        if not lemma:
            continue
        leading = split_leading(definition_text)
        reading = clean(cells[2])
        yield _row(
            id=f"2569-{index:05d}", lemma=lemma, pos=clean(cells[3]) or leading.pos,
            sense_number=clean(cells[1]).strip("()"),
            definition=leading.remainder or definition_text,
            pronunciations=([reading] if reading and reading != lemma else []) + leading.pronunciations,
            notes=leading.labels, raw=_raw(header, cells),
        )


def _technical_term(text: str) -> tuple[list[str], str, list[str]]:
    """Thai equivalents, POS and notes from one ศัพท์บัญญัติ segment."""
    leading = split_leading(text, loose_pos=True)
    notes = list(leading.labels)
    term = leading.remainder
    for bracket in re.findall(r"\[([^\]]+)\]", term):
        notes.append(f"[{bracket.strip()}]")
    term = re.sub(r"\[[^\]]+\]", "", term)
    trailing = re.search(r"\s*\(([^()]+)\)\s*$", term)
    if trailing:
        notes.append(f"({trailing.group(1).strip()})")
        term = term[: trailing.start()]
    alternatives = [part.strip(" .;") for part in re.split(r"[,;]", term)]
    return [item for item in alternatives if item], leading.pos, notes


def technical_terms(header: list[str], rows: Iterable[list[object]]) -> Iterator[dict[str, str]]:
    """Royal Institute technical glossaries: English term → Thai term(s) + Thai explanation."""
    english = ""
    for cells in rows:
        cells = list(cells) + [""] * (9 - len(cells))
        number, thai = clean(cells[0]), clean(cells[2])
        english = clean(cells[1]) or english
        if not number or not thai or not english:
            continue
        description = clean(cells[8])
        described = dict(numbered_segments(description)) if description else {}
        segments = numbered_segments(thai)
        common_notes = [
            f"{label}: {value}" for label, value in (
                ("สาขาวิชา", clean(cells[3])),
                ("ปีที่พิมพ์", clean(cells[4]).removesuffix(" 00:00:00")),
                ("ศัพท์ตั้งอื่น", clean(cells[5])),
                ("คำอ้างอิง", clean(cells[6])),
            ) if value
        ]
        for segment_number, segment in segments:
            terms, pos, notes = _technical_term(segment)
            definition = described.get(segment_number) if segment_number and len(described) > 1 else description
            for alternative, term in enumerate(terms, start=1):
                yield _row(
                    id=f"{number}-{segment_number or 0}-{alternative}", lemma=term, pos=pos,
                    sense_number=segment_number, definition=definition or "",
                    notes=[*notes, *common_notes], translations=[english],
                    raw=_raw(header, cells),
                )


def transliterations(header: list[str], rows: Iterable[list[object]]) -> Iterator[dict[str, str]]:
    """Columns: ลำดับที่, English word, Thai transliteration, note."""
    for cells in rows:
        cells = list(cells) + [""] * (4 - len(cells))
        number, english, thai, note = (clean(value) for value in cells[:4])
        if not number or not english or not thai:
            continue
        for alternative, term in enumerate(re.split(r",|\sหรือ\s", thai), start=1):
            if term.strip():
                yield _row(
                    id=f"{number}-{alternative}", lemma=term.strip(), translations=[english],
                    notes=[f"หมายเหตุ: {note}"] if note and note != "-" else [],
                    raw=_raw(header, cells),
                )


# --- Word profile ---------------------------------------------------------

_DIALECT_BRACKETS = re.compile(r"\[[^\]]+\]")


def dialect_paragraphs(paragraphs: Iterable[str], prefix: str) -> Iterator[dict[str, str]]:
    """Dialect glossaries: `headword<TAB>script<TAB>[Thai reading]<TAB>[IPA]` then the definition.

    The regional-script columns rely on a legacy font and are kept only in the
    raw record.
    """
    for index, paragraph in enumerate(paragraphs, start=1):
        if "\t" not in paragraph:
            continue
        lines = [line for line in paragraph.replace("\r", "\n").split("\n") if line.strip()]
        header_parts = lines[0].split("\t")
        lemma = clean(header_parts[0])
        if not lemma:
            continue
        header_text = "\t".join(header_parts[1:])
        body = lines[1:]
        while body and body[0].strip().startswith("["):
            header_text += "\t" + body.pop(0)
        brackets = _DIALECT_BRACKETS.findall(header_text)
        definition_text = clean(" ".join(body))
        if not definition_text:
            last = max((header_text.rfind(item) + len(item) for item in brackets), default=0)
            definition_text = clean(header_text[last:])
        if not definition_text:
            continue
        number = re.match(r"^([๐-๙0-9]{1,2})\.\s+", definition_text)
        leading = split_leading(definition_text[number.end():] if number else definition_text)
        yield _row(
            id=f"{prefix}-{index:05d}", lemma=lemma, pos=leading.pos,
            definition=leading.remainder or definition_text,
            pronunciations=[item.strip("[] ") for item in brackets] + leading.pronunciations,
            notes=leading.labels, raw=json.dumps({"paragraph": paragraph}, ensure_ascii=False),
        )


# --- IO -------------------------------------------------------------------

SPREADSHEET_PROFILES = {
    "royal-2542": royal_2542,
    "royal-2554": royal_2554,
    "royal-2569": royal_2569,
    "technical-terms": technical_terms,
    "transliterations": transliterations,
}
PROFILES = (*SPREADSHEET_PROFILES, "dialect-docx")


def _read_spreadsheet(path: Path) -> tuple[list[str], list[list[object]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    header = [clean(value) for value in rows[0]] if rows else []
    return header, rows[1:]


def prepare(profile: str, input_path: Path, output_path: Path) -> dict[str, object]:
    if profile in SPREADSHEET_PROFILES:
        header, rows = _read_spreadsheet(input_path)
        prepared = list(SPREADSHEET_PROFILES[profile](header, rows))
        source_rows = len(rows)
    elif profile == "dialect-docx":
        import docx

        paragraphs = [paragraph.text for paragraph in docx.Document(str(input_path)).paragraphs]
        prepared = list(dialect_paragraphs(paragraphs, output_path.stem))
        source_rows = len([item for item in paragraphs if item.strip()])
    else:
        raise ValueError(f"Unknown profile {profile!r}; expected one of {PROFILES}")

    seen: set[str] = set()
    for row in prepared:
        if row["id"] in seen:
            raise ValueError(f"Duplicate prepared id {row['id']!r}")
        seen.add(row["id"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(prepared)
    return {
        "profile": profile,
        "input": str(input_path),
        "output": str(output_path),
        "source_rows": source_rows,
        "prepared_rows": len(prepared),
        "with_definition": sum(bool(row["definition"]) for row in prepared),
        "with_pos": sum(bool(row["pos"]) for row in prepared),
        "unique_lemmas": len({row["lemma"] for row in prepared}),
    }
